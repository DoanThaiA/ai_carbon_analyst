"""Trích chữ + chia đoạn tài liệu tri thức (services/knowledge_extraction.py). PDF test được dựng tay bằng
_make_pdf (font Helvetica chuẩn, chỉ chữ ASCII) để không phải commit file nhị phân."""
import io
import unicodedata
import zipfile

import pytest
from docx import Document

from services import knowledge_extraction as kx
from services.knowledge_extraction import (
    ExtractedSection,
    KnowledgeExtractionError,
    chunk_document,
    detect_media_type,
    embedding_text,
    extract_docx,
    extract_document,
    extract_pdf,
    join_pdf_lines,
    normalize_text,
)


def _esc(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _make_pdf(pages: list) -> bytes:
    """pages: [{"lines": [(x, y, text)], "rects": [(x, y, w, h)], "image": bool}] → bytes PDF 1.4 tối giản."""
    objects: list = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    image = add(
        b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray /BitsPerComponent 8 "
        b"/Length 1 >>\nstream\n\x80\nendstream"
    )
    pages_id = add(b"")
    kids = []
    for p in pages:
        ops = [f"{x} {y} {w} {h} re S" for x, y, w, h in p.get("rects", [])]
        ops += [f"BT /F1 10 Tf {x} {y} Td ({_esc(t)}) Tj ET" for x, y, t in p.get("lines", [])]
        if p.get("image"):
            ops.append("q 300 0 0 300 100 300 cm /Im1 Do Q")
        stream = "\n".join(ops).encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        kids.append(add(
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] /Contents {content} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> /XObject << /Im1 {image} 0 R >> >> >>".encode()
        ))
    objects[pages_id - 1] = (
        f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>".encode()
    )
    catalog = add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode())
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def _text_page(no: int, body: list) -> dict:
    """Trang có header/footer lặp (tên tài liệu + số trang) và các dòng thân bài từ trên xuống."""
    lines = [(72, 760, "EU ETS Handbook 2026"), (300, 40, str(no))]
    lines += [(72, 700 - 14 * i, text) for i, text in enumerate(body)]
    return {"lines": lines}


def _docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── Chuẩn hoá + nhận diện định dạng ──────────────────────────────────

def test_normalize_text_nfc_invisible_and_spaces():
    decomposed = unicodedata.normalize("NFD", "Hạn ngạch phát thải")
    text = normalize_text(f"  {decomposed}­  ﬁnal  số   liệu\n\n\n\nđoạn 2 ")
    assert text == "Hạn ngạch phát thải final số liệu\n\nđoạn 2"
    assert unicodedata.is_normalized("NFC", text)


def test_detect_media_type_by_magic_bytes():
    assert detect_media_type(_make_pdf([_text_page(1, ["Hello"])])) == "pdf"
    assert detect_media_type(_docx_bytes(lambda d: d.add_paragraph("x"))) == "docx"
    with pytest.raises(KnowledgeExtractionError, match=r"\.docx"):
        detect_media_type(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
    with pytest.raises(KnowledgeExtractionError):
        detect_media_type(b"not a document at all")


def test_extract_document_rejects_oversized(monkeypatch):
    monkeypatch.setattr(kx, "KNOWLEDGE_MAX_BYTES", 100)
    with pytest.raises(KnowledgeExtractionError, match="MB"):
        extract_document(b"%PDF-" + b"0" * 200)


# ── DOCX ─────────────────────────────────────────────────────────────

def test_docx_keeps_order_headings_and_tables():
    def build(d):
        d.core_properties.title = "Quy định CBAM"
        d.add_paragraph("Lời mở đầu của tài liệu.")
        d.add_heading("Điều 1. Phạm vi", level=1)
        d.add_paragraph("Áp dụng cho thép và nhôm nhập khẩu vào EU.")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Mã CN", "Benchmark"
        table.cell(1, 0).text, table.cell(1, 1).text = "7601", "1,423"
        d.add_paragraph("Đoạn sau bảng.")
        d.add_heading("Điều 2. Hiệu lực", level=1)
        d.add_paragraph("Từ 01/01/2026.")

    doc = extract_docx(_docx_bytes(build))
    assert [s.heading for s in doc.sections] == [None, "Điều 1. Phạm vi", "Điều 2. Hiệu lực"]
    body = doc.sections[1].text
    assert body.index("Áp dụng cho thép") < body.index("Mã CN | Benchmark") < body.index("Đoạn sau bảng.")
    assert "7601 | 1,423" in body
    assert doc.title_hint == "Quy định CBAM"
    assert "## Điều 2. Hiệu lực" in doc.content


def test_docx_merged_cells_not_repeated():
    def build(d):
        table = d.add_table(rows=2, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 2))
        merged.text = "Tiêu đề gộp"
        for i, v in enumerate(("a", "b", "c")):
            table.cell(1, i).text = v

    text = extract_docx(_docx_bytes(build)).sections[0].text
    assert text.split("\n") == ["Tiêu đề gộp", "a | b | c"]


def test_docx_empty_raises():
    with pytest.raises(KnowledgeExtractionError, match="Không trích được"):
        extract_docx(_docx_bytes(lambda d: d.add_paragraph("   ")))


def test_docx_zip_bomb_guard(monkeypatch):
    monkeypatch.setattr(kx, "MAX_DOCX_UNCOMPRESSED_BYTES", 1000)
    data = _docx_bytes(lambda d: d.add_paragraph("x" * 5000))
    with pytest.raises(KnowledgeExtractionError, match="giải nén"):
        extract_docx(data)


def test_docx_with_macro_rejected():
    source = zipfile.ZipFile(io.BytesIO(_docx_bytes(lambda d: d.add_paragraph("x"))))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as out:
        for item in source.infolist():
            out.writestr(item, source.read(item))
        out.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(KnowledgeExtractionError, match="macro"):
        extract_docx(buffer.getvalue())


# ── PDF ──────────────────────────────────────────────────────────────

def test_join_pdf_lines_paragraphs_lists_and_hyphenation():
    lines = [
        "The Market Stability Reserve absorbs surplus allowances when the total number in circula-",
        "tion exceeds the upper threshold set by the regulation.",
        "Short closing line.",
        "- First bullet item",
        "- Second bullet item",
        "Article 3 Scope of application",
    ]
    text = join_pdf_lines(lines)
    paragraphs = text.split("\n\n")
    assert paragraphs[0].startswith("The Market Stability Reserve") and "circulation exceeds" in paragraphs[0]
    assert "- First bullet item" in paragraphs and "- Second bullet item" in paragraphs
    assert paragraphs[-1] == "Article 3 Scope of application"


def test_pdf_text_pages_drop_header_footer_and_page_numbers():
    topics = ["the cap declines by 4.3 percent", "auctions follow the calendar", "the reserve absorbs surplus", "free allocation phases out"]
    pages = [
        _text_page(no, [f"Chapter {no} explains how {topic} each year.", f"Body line {no * 7} continues here."])
        for no, topic in enumerate(topics, start=1)
    ]
    doc = extract_pdf(_make_pdf(pages))
    assert doc.page_count == 4
    assert [s.page_from for s in doc.sections] == [1, 2, 3, 4]
    assert "EU ETS Handbook" not in doc.content
    assert "Chapter 3 explains how the reserve absorbs surplus each year." in doc.sections[2].text
    assert all(s.text.strip() not in {"1", "2", "3", "4"} for s in doc.sections)
    assert not any(line.strip().isdigit() for s in doc.sections for line in s.text.split("\n"))


def test_pdf_table_extracted_as_rows_without_duplicating_text():
    grid = [(100 + 100 * c, 600 + 30 * r, 100, 30) for c in range(2) for r in range(3)]
    cells = [
        (105, 665, "Product"), (205, 665, "Benchmark"),
        (105, 635, "Aluminium"), (205, 635, "1.423"),
        (105, 605, "Steel"), (205, 605, "1.370"),
    ]
    page = {"rects": grid, "lines": [(72, 720, "Default benchmarks per tonne of product are listed below.")] + cells}
    doc = extract_pdf(_make_pdf([page]))
    text = doc.sections[0].text
    assert "Product | Benchmark" in text and "Aluminium | 1.423" in text and "Steel | 1.370" in text
    assert text.count("1.423") == 1


def test_pdf_mostly_scanned_rejected():
    with pytest.raises(KnowledgeExtractionError, match="scan"):
        extract_pdf(_make_pdf([{"image": True}, {"image": True}, _text_page(3, ["Some real text on this page here."])]))


def test_pdf_some_scanned_pages_warn():
    pages = [_text_page(no, [f"Page {no} has enough readable text for extraction."]) for no in range(1, 4)]
    pages.append({"image": True})
    doc = extract_pdf(_make_pdf(pages))
    assert len(doc.sections) == 3
    assert doc.warnings and "scan" in doc.warnings[0] and "4" in doc.warnings[0]


def test_pdf_too_many_pages(monkeypatch):
    monkeypatch.setattr(kx, "MAX_PAGES", 2)
    with pytest.raises(KnowledgeExtractionError, match="giới hạn"):
        extract_pdf(_make_pdf([_text_page(n, ["text"]) for n in range(1, 4)]))


def test_pdf_corrupt_rejected():
    with pytest.raises(KnowledgeExtractionError):
        extract_pdf(b"%PDF-1.4\ngarbage")


# ── Chia đoạn ────────────────────────────────────────────────────────

def test_chunk_document_respects_headings_and_page_ranges():
    sections = [
        ExtractedSection(heading=None, page_from=1, page_to=1, text="A" * 400),
        ExtractedSection(heading=None, page_from=2, page_to=2, text="B" * 400),
        ExtractedSection(heading="Điều 2", page_from=None, page_to=None, text="C" * 300),
    ]
    chunks = chunk_document(sections, max_chars=1000)
    assert len(chunks) == 2
    assert (chunks[0].page_from, chunks[0].page_to, chunks[0].heading) == (1, 2, None)
    assert chunks[1].heading == "Điều 2" and chunks[1].text == "C" * 300
    assert chunks[0].meta == {"heading": None, "page_from": 1, "page_to": 2}


def test_chunk_document_splits_long_table_with_header():
    rows = ["Mã CN | Benchmark"] + [f"{7000 + i} | 1,{i:03d} tCO2/t" for i in range(80)]
    sections = [ExtractedSection(heading="Phụ lục", page_from=5, page_to=5, text="\n".join(rows))]
    chunks = chunk_document(sections, max_chars=400)
    assert len(chunks) > 1
    assert all(c.text.startswith("Mã CN | Benchmark") for c in chunks)
    assert all(len(c.text) <= 400 for c in chunks)
    body_rows = [line for c in chunks for line in c.text.split("\n")[1:]]
    assert body_rows == rows[1:]


def test_chunk_document_splits_long_prose_by_sentence():
    text = " ".join(f"Câu số {i} nói về giá EUA và cơ chế MSR." for i in range(100))
    chunks = chunk_document([ExtractedSection(heading=None, page_from=1, page_to=1, text=text)], max_chars=300)
    assert len(chunks) > 1 and all(len(c.text) <= 300 for c in chunks)


def test_embedding_text_prefix():
    chunk = chunk_document([ExtractedSection(heading="Điều 1", page_from=None, page_to=None, text="Nội dung.")])[0]
    assert embedding_text("Quy định CBAM", chunk) == "Tài liệu: Quy định CBAM — Điều 1\nNội dung."
