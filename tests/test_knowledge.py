"""Tri thức tài liệu cho Jenny (services/knowledge.py): phạm vi dùng, khối prompt + mã [TLk], trích dẫn trong báo
cáo (report_generator), Jack đối chiếu số với đoạn tài liệu (report_qc), tool chat không lộ ra MCP."""
import io
from datetime import date
from types import SimpleNamespace

import pytest
from docx import Document

from schemas.retrieval_models import KnowledgeHit
from services import knowledge
from services import quote_chat as qc
from services import report_generator as rg
from services import report_qc


def _hit(i: int = 1, doc_id: int = 7, **kw) -> KnowledgeHit:
    base = dict(
        chunk_id=100 + i, doc_id=doc_id, chunk_index=i, content=f"Đoạn {i}: hệ số CBAM năm 2027 là 97,5% và benchmark 1,423.",
        score=0.8, title="Quy định CBAM 2025/2620", publisher="Ủy ban châu Âu", issued_on=date(2025, 12, 1),
        verified=True, page_from=12, page_to=12,
    )
    base.update(kw)
    return KnowledgeHit(**base)


def _doc(**kw):
    base = dict(
        status="ready", is_active=True, valid_from=date(2026, 1, 1), valid_to=None, use_in_chat=True,
        use_in_report=True, report_sections=["3", "4"],
    )
    base.update(kw)
    return SimpleNamespace(**base)


# ── Phạm vi dùng ──────────────────────────────────────────────────────

def test_is_applicable_scope_and_validity():
    as_of = date(2026, 10, 21)
    assert knowledge.is_applicable(_doc(), as_of, "report", "3")
    assert not knowledge.is_applicable(_doc(), as_of, "report", "1")  # mục không được chọn
    assert knowledge.is_applicable(_doc(report_sections=None), as_of, "report", "1")  # NULL = mọi mục
    assert knowledge.is_applicable(_doc(), as_of, "report")  # không hỏi mục cụ thể
    assert knowledge.is_applicable(_doc(), as_of, "chat")
    assert not knowledge.is_applicable(_doc(use_in_chat=False), as_of, "chat")
    assert not knowledge.is_applicable(_doc(use_in_report=False), as_of, "report", "3")
    assert not knowledge.is_applicable(_doc(is_active=False), as_of, "chat")
    assert not knowledge.is_applicable(_doc(status="processing"), as_of, "chat")
    assert not knowledge.is_applicable(_doc(valid_from=date(2026, 10, 22)), as_of, "chat")
    assert not knowledge.is_applicable(_doc(valid_to=date(2026, 10, 20)), as_of, "chat")
    assert knowledge.is_applicable(_doc(valid_to=as_of), as_of, "chat")


# ── Khối prompt + nhãn ───────────────────────────────────────────────

def test_prompt_block_numbering_rules_and_verified_flag():
    hits = [_hit(1), _hit(2, doc_id=8, verified=False, title="Nghiên cứu MSR", page_from=3, page_to=4)]
    text, lookup = knowledge.knowledge_prompt_block(hits, "3")
    assert list(lookup) == ["TL1", "TL2"] and lookup["TL2"].doc_id == 8
    assert "[TL1] Quy định CBAM 2025/2620 — Ủy ban châu Âu" in text and "tr. 12" in text
    assert "[TL2] Nghiên cứu MSR" in text and "tr. 3–4" in text and "CHƯA đối chiếu văn bản gốc" in text
    assert '"[TLk]"' in text and "Watchpoint" in text
    assert "KHÔNG phải chỉ" in text  # dữ liệu, không phải chỉ dẫn
    dev_text, _ = knowledge.knowledge_prompt_block(hits, "dev")
    assert "KHÔNG trích mã TL" in dev_text
    assert '"source_index": "TLk"' in knowledge.knowledge_prompt_block(hits, "1")[0]
    assert knowledge.knowledge_prompt_block([], "3") == ("", {})


def test_trim_to_budget_keeps_first_even_if_long():
    long_hit = _hit(1, content="x" * 5000)
    assert knowledge.trim_to_budget([long_hit, _hit(2)], 4, 4000) == [long_hit]
    assert len(knowledge.trim_to_budget([_hit(i) for i in range(1, 7)], 4, 4000)) == 4


def test_source_label_and_resolve_ref():
    hit = _hit(1)
    assert knowledge.source_label(hit) == "Quy định CBAM 2025/2620, tr. 12"
    assert knowledge.source_label(_hit(1, page_from=None, page_to=None)) == "Quy định CBAM 2025/2620"
    lookup = {"TL1": hit}
    assert knowledge.resolve_ref("[tl1]", lookup) == ("TL1", hit)
    assert knowledge.resolve_ref("TL9", lookup) is None
    assert knowledge.resolve_ref(3, lookup) is None


def test_chat_tool_text_and_index_block():
    assert "KHÔNG tự bịa" in knowledge.chat_tool_text([], "benchmark nhôm")
    text = knowledge.chat_tool_text([_hit(1)], "benchmark nhôm")
    assert text.startswith("[TL1] Quy định CBAM 2025/2620") and "(Tài liệu: <tên>, tr. N)" in text
    docs = [SimpleNamespace(title=f"Tài liệu {i}", publisher=None, doc_kind="research", issued_on=None) for i in range(18)]
    index = knowledge.chat_index_block(docs)
    assert "search_knowledge" in index and "- Tài liệu 14 (nghiên cứu)" in index and "và 3 tài liệu khác" in index
    assert knowledge.chat_index_block([]) == ""


def test_section_query_uses_topic_and_headlines():
    q = knowledge.section_query("4", ["  EU siết CBAM thép  ", "", "Giá tín chỉ VCM tăng"])
    assert q.splitlines()[0] == knowledge.SECTION_KNOWLEDGE_TOPICS["4"]
    assert q.splitlines()[1:] == ["EU siết CBAM thép", "Giá tín chỉ VCM tăng"]


# ── Trích dẫn trong báo cáo ──────────────────────────────────────────

def test_section3_citations_resolve_tl_and_drop_tl_only_watchpoint():
    hit = _hit(1)
    sec3 = {"analysis_blocks": [
        {"heading": "Phân tích", "content": "Hệ số CBAM năm 2027 là 97,5% [TL1]; tin hôm nay [1]. Mã lạ [TL9]."},
        {"heading": rg.WATCH_HEADING, "content": "1. Chỉ dựa tài liệu → theo dõi [TL1]\n2. Tin nguồn cung [1]"},
    ]}
    index = {1: {"url": "https://a.example/x", "source": "a.example"}}
    sources = rg._resolve_section3_citations(sec3, index, {}, {"TL1": hit})
    analysis = sec3["analysis_blocks"][0]["content"]
    assert "(Nguồn: Quy định CBAM 2025/2620, tr. 12" in analysis and "[TL" not in analysis
    watch = sec3["analysis_blocks"][1]["content"]
    assert watch.startswith("1. Tin nguồn cung") and "Chỉ dựa tài liệu" not in watch
    kb = [s for s in sources if s.get("source_type") == "knowledge"]
    assert kb == [{"source_name": "Quy định CBAM 2025/2620", "source_type": "knowledge", "doc_id": 7, "chunk_ids": [101]}]


def test_bullet_sources_accept_tl_index_and_inline_mark():
    hit = _hit(2)
    bullets = [
        {"text": "**CBAM:** hệ số 97,5% năm 2027.", "source_index": "TL1"},
        {"text": "**CBAM:** benchmark nhôm 1,423 [TL1].", "source_index": None},
        {"text": "**EUA:** tin hôm nay.", "source_index": 1},
    ]
    out = rg._resolve_bullet_sources(bullets, {1: {"url": "https://a.example/x", "source": "a.example"}}, {"TL1": hit})
    assert out[0]["source_url"] is None and out[0]["source_name"] == "Quy định CBAM 2025/2620, tr. 12"
    assert out[0]["knowledge_ref"] == {"code": "TL1", "doc_id": 7, "chunk_id": 102, "title": "Quy định CBAM 2025/2620"}
    assert out[1]["text"] == "**CBAM:** benchmark nhôm 1,423." and out[1]["knowledge_ref"]["chunk_id"] == 102
    assert out[2]["source_url"] == "https://a.example/x" and "knowledge_ref" not in out[2]
    # Không có tài liệu → giữ hành vi cũ (không nguồn).
    assert rg._resolve_bullet_sources([{"text": "x", "source_index": "TL1"}], {})[0]["source_name"] is None


def test_cited_documents_and_section9_items():
    content = {
        "1": {"bullets": [{"text": "a", "knowledge_ref": {"doc_id": 7, "chunk_id": 101, "title": "Quy định CBAM"}}]},
        "3": {"sources": [{"source_type": "knowledge", "doc_id": 8, "source_name": "Nghiên cứu MSR", "title": "Nghiên cứu MSR"},
                          {"source_type": "knowledge", "doc_id": 7, "source_name": "Quy định CBAM"}]},
        "biz": {"long_term": [{"opportunity": "x", "knowledge_ref": {"doc_id": 9, "chunk_id": 1, "title": "Nội bộ SIM"}}]},
    }
    assert [d["doc_id"] for d in knowledge.cited_documents(content)] == [7, 8, 9]
    items = rg._knowledge_section9_items(content, {"3": [_hit(1)]})
    assert items[0] == {"source": "Ủy ban châu Âu", "title": "Quy định CBAM", "url": None, "kind": "knowledge", "doc_id": 7}
    assert items[2]["source"] == "Tài liệu tham khảo"


# ── Jack QC ──────────────────────────────────────────────────────────

def test_check_facts_uses_knowledge_chunk_text():
    content = {
        "1": {"bullets": [
            {"text": "**CBAM:** hệ số 97,5% năm 2027, benchmark 1,423.", "source_name": "Quy định CBAM, tr. 12",
             "knowledge_ref": {"code": "TL1", "doc_id": 7, "chunk_id": 101, "title": "Quy định CBAM"}},
            {"text": "**CBAM:** tỷ lệ miễn phí 45,6%.", "source_name": "Quy định CBAM, tr. 12",
             "knowledge_ref": {"code": "TL1", "doc_id": 7, "chunk_id": 101, "title": "Quy định CBAM"}},
        ]},
        "3": {
            "sources": [{"source_type": "knowledge", "doc_id": 7, "source_name": "Quy định CBAM", "chunk_ids": [101]}],
            "analysis_blocks": [{"heading": "Phân tích", "content": "Hệ số 2027 là 88,8% (Nguồn: Quy định CBAM, tr. 12)."}],
        },
    }
    texts = {101: "Hệ số CBAM năm 2027 là 97,5%; benchmark nhôm 1,423 tCO2/t."}
    issues = report_qc.check_facts(content, {}, texts)
    messages = [i["message"] for i in issues if i["check"] == "fact"]
    assert not any("bullet 1" in m for m in messages)
    assert any("bullet 2" in m and "45,6" in m for m in messages)
    assert any("88,8" in m for m in messages)
    # Không có chữ đoạn (tài liệu đã xoá) → bỏ qua, không báo sai.
    assert not [i for i in report_qc.check_facts(content, {}, {}) if "không có trong bài nguồn" in i["message"]]


def test_check_sources_skips_knowledge_items_and_refs_state():
    content = {"9": {"items": [
        {"source": "Ủy ban châu Âu", "title": "Quy định CBAM", "url": None, "kind": "knowledge", "doc_id": 7},
    ]}, "1": {"bullets": [{"text": "a", "knowledge_ref": {"doc_id": 7, "chunk_id": 1, "title": "Quy định CBAM"}}]}}
    assert not report_qc.check_sources(content, [])
    assert report_qc.check_knowledge_refs(content, {7: "active"}) == []
    assert report_qc.check_knowledge_refs(content, {7: "inactive"})[0]["severity"] == "info"
    deleted = report_qc.check_knowledge_refs(content, {})
    assert deleted[0]["severity"] == "warning" and "đã bị xoá" in deleted[0]["message"]
    assert report_qc.knowledge_chunk_ids(
        {"knowledge_used": {"3": [{"chunk_id": 5}]}, "3": {"sources": [{"source_type": "knowledge", "chunk_ids": [6]}]}}
    ) == {5, 6}


def test_translation_skips_knowledge_subtrees():
    from services import report_translation as rt

    out: list = []
    rt._collect({"knowledge_used": {"3": [{"title": "Tài liệu tiếng Việt"}]},
                 "1": {"bullets": [{"text": "Nội dung tiếng Việt", "knowledge_ref": {"title": "Tên tài liệu có dấu"}}]}}, None, out)
    assert "Nội dung tiếng Việt" in " ".join(out)
    assert not any("tài liệu" in t.lower() for t in out)


# ── Jenny chat / MCP ─────────────────────────────────────────────────

def test_knowledge_tool_never_exposed_to_mcp():
    """Tài liệu có thể là báo cáo trả phí/nội bộ — gateway Claude Desktop/mobile không phục vụ."""
    from api.routers import mcp_gateway

    assert qc.KNOWLEDGE_TOOL["name"] not in {t["name"] for t in qc.CLIENT_TOOLS}
    assert qc.KNOWLEDGE_TOOL["name"] not in mcp_gateway._TOOLS_BY_NAME


def test_static_prompt_includes_knowledge_index_only_when_given():
    block = knowledge.chat_index_block([SimpleNamespace(title="Quy định CBAM", publisher=None, doc_kind="regulation", issued_on=None)])
    assert "Quy định CBAM" in qc._build_static_instructions("2026-10-15", None, "", block)
    assert "TÀI LIỆU THAM KHẢO DO ADMIN" not in qc._build_static_instructions("2026-10-15")


# ── Upload / schema ──────────────────────────────────────────────────

def test_validate_knowledge_upload():
    from services import minio_service as ms

    ms.validate_knowledge_upload("a.pdf", ms.PDF_MEDIA_TYPE, 1000)
    ms.validate_knowledge_upload("A.DOCX", ms.DOCX_MEDIA_TYPE, 1000)
    with pytest.raises(ms.InvalidAttachmentError, match=r"\.doc"):
        ms.validate_knowledge_upload("a.doc", ms.LEGACY_DOC_MEDIA_TYPE, 1000)
    with pytest.raises(ms.InvalidAttachmentError):
        ms.validate_knowledge_upload("a.pdf", ms.DOCX_MEDIA_TYPE, 1000)  # đuôi ≠ loại khai báo
    with pytest.raises(ms.InvalidAttachmentError):
        ms.validate_knowledge_upload("a.png", "image/png", 1000)
    with pytest.raises(ms.InvalidAttachmentError, match="MB"):
        ms.validate_knowledge_upload("a.pdf", ms.PDF_MEDIA_TYPE, ms.KNOWLEDGE_MAX_BYTES + 1)


def test_create_request_validation():
    from pydantic import ValidationError

    from schemas.knowledge_models import KnowledgeCreateRequest, KnowledgeUpdateRequest

    req = KnowledgeCreateRequest(file_key="knowledge/x.pdf", file_name="x.pdf", title="  Quy định  ", report_sections=["biz", "3"])
    assert req.title == "Quy định" and req.report_sections == ["3", "biz"]
    assert KnowledgeUpdateRequest(report_sections=[]).report_sections is None
    with pytest.raises(ValidationError):
        KnowledgeCreateRequest(file_key="k", file_name="x.pdf", title="t", report_sections=["9"])
    with pytest.raises(ValidationError):
        KnowledgeCreateRequest(file_key="k", file_name="x.pdf", title="t", valid_from=date(2026, 5, 1), valid_to=date(2026, 4, 1))


def test_extract_isolated_runs_in_subprocess():
    document = Document()
    document.add_heading("Điều 1", level=1)
    document.add_paragraph("Benchmark nhôm 1,423 tCO2/t.")
    buffer = io.BytesIO()
    document.save(buffer)
    doc = knowledge._extract_isolated(buffer.getvalue(), 60)
    assert doc.sections[0].heading == "Điều 1" and "1,423" in doc.sections[0].text
    with pytest.raises(knowledge.KnowledgeExtractionError):
        knowledge._extract_isolated(b"not a document", 60)


# ── Truy hồi cho báo cáo (session/retrieval giả, không cần Postgres/Cohere) ──

class _FakeSession:
    def __init__(self, docs):
        self.docs = docs
        self.savepoints = 0

    async def scalars(self, _stmt):
        return list(self.docs)

    def begin_nested(self):
        session = self

        class _Savepoint:
            async def __aenter__(self):
                session.savepoints += 1

            async def __aexit__(self, *exc):
                return False  # không nuốt lỗi — giống SAVEPOINT thật (rollback rồi raise lại)

        return _Savepoint()


class _FakeRetrieval:
    def __init__(self, hits_by_section, fail=False):
        self.hits_by_section = hits_by_section
        self.fail = fail
        self.calls = []
        self.embed_batches = []
        outer = self

        class _Embedder:
            async def embed(self, texts, input_type="search_document"):
                outer.embed_batches.append((list(texts), input_type))
                return [[0.0] for _ in texts]

        self.embedder = _Embedder()

    async def retrieve_knowledge(self, query, *, purpose, as_of, section=None, **kw):
        if self.fail:
            raise RuntimeError("cohere down")
        self.calls.append((section, purpose, as_of, kw["top_k"]))
        return self.hits_by_section.get(section, [])


def _run(coro):
    import asyncio

    return asyncio.run(coro)


def test_load_report_knowledge_no_documents_makes_no_calls():
    created = []
    out = _run(knowledge.load_report_knowledge(
        _FakeSession([]), lambda: created.append(1), "2026-10-21", {}, rg.SECTION_TOPICS, ["1", "3"],
    ))
    assert out == {} and created == []


def test_load_report_knowledge_batches_queries_and_filters_sections():
    news = {"cbam": [{"title": "EU siết CBAM thép"}], "eua_ets": [{"title": "EUA giảm 2%"}]}
    fake = _FakeRetrieval({"3": [_hit(1)], "4": [_hit(2), _hit(3, content="y" * 5000)]})
    out = _run(knowledge.load_report_knowledge(
        _FakeSession([_doc(report_sections=["3", "4"])]), lambda: fake, "2026-10-21", news, rg.SECTION_TOPICS,
        ["1", "3", "4"],
    ))
    assert [len(batch) for batch, _ in fake.embed_batches] == [2] and fake.embed_batches[0][1] == "search_query"
    assert "EU siết CBAM thép" in fake.embed_batches[0][0][1]
    assert [c[0] for c in fake.calls] == ["3", "4"] and all(c[1] == "report" for c in fake.calls)
    assert fake.calls[0][2] == date(2026, 10, 21)
    assert [h.chunk_id for h in out["3"]] == [101]
    assert [h.chunk_id for h in out["4"]] == [102]  # đoạn 5000 ký tự vượt ngân sách sau đoạn đầu


def test_load_report_knowledge_errors_never_break_report():
    out = _run(knowledge.load_report_knowledge(
        _FakeSession([_doc()]), lambda: _FakeRetrieval({}, fail=True), "2026-10-21", {}, rg.SECTION_TOPICS, ["3"],
    ))
    assert out == {}


def test_chat_search_text_clamps_input_and_handles_errors():
    fake = _FakeRetrieval({None: [_hit(1)]})
    assert _run(knowledge.chat_search_text(fake, {"query": ""})) == "Cần truyền query."
    text = _run(knowledge.chat_search_text(fake, {"query": "benchmark nhôm", "top_k": 99}))
    assert text.startswith("[TL1]") and fake.calls[0][1] == "chat" and fake.calls[0][3] == knowledge.CHAT_MAX_HITS
    assert "KHÔNG tự bịa" in _run(knowledge.chat_search_text(_FakeRetrieval({}, fail=True), {"query": "x"}))


# ── Job xử lý tài liệu (session/MinIO/embedder giả) ──────────────────

class _ProcSession:
    def __init__(self, doc):
        self.doc = doc
        self.executed = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    async def get(self, _model, _id):
        return self.doc

    async def commit(self):
        pass

    async def refresh(self, _obj):
        pass

    async def rollback(self):
        pass

    async def execute(self, stmt):
        self.executed.append(stmt)


class _FakeEmbedder:
    def __init__(self):
        self.batches = []

    async def embed(self, texts, input_type="search_document"):
        self.batches.append((len(texts), input_type))
        return [[0.1, 0.2] for _ in texts]


def _docx_with_paragraphs(n: int) -> bytes:
    document = Document()
    document.add_heading("Phụ lục I", level=1)
    for i in range(n):
        document.add_paragraph(f"Đoạn {i}: benchmark sản phẩm {i} là 1,{i:03d} tCO2/t theo quy định. " * 4)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_process_document_indexes_chunks_in_batches(monkeypatch):
    from services import minio_service
    from services.knowledge_extraction import extract_document

    monkeypatch.setattr(minio_service, "get_file_as_bytes", lambda key: _docx_with_paragraphs(300))
    monkeypatch.setattr(knowledge, "EMBED_BATCH_SIZE", 50)
    doc = SimpleNamespace(id=5, status="pending", title="Phụ lục CBAM", file_key="knowledge/x.docx", error_message=None)
    session, embedder = _ProcSession(doc), _FakeEmbedder()
    _run(knowledge.process_document(lambda: session, embedder, 5, extractor=lambda data, _t: extract_document(data)))
    assert doc.status == "ready" and doc.chunk_count > 50 and doc.error_message is None
    assert all(n <= 50 for n, _ in embedder.batches) and sum(n for n, _ in embedder.batches) == doc.chunk_count
    assert "## Phụ lục I" in doc.content and doc.sections[0]["heading"] == "Phụ lục I"
    assert [type(s).__name__ for s in session.executed] == ["Delete", "Insert"]


def test_process_document_marks_failed_on_extraction_error(monkeypatch):
    from services import minio_service

    monkeypatch.setattr(minio_service, "get_file_as_bytes", lambda key: b"garbage")
    doc = SimpleNamespace(id=6, status="pending", title="Hỏng", file_key="knowledge/y.pdf", error_message=None)
    session = _ProcSession(doc)

    def _boom(data, _t):
        raise knowledge.KnowledgeExtractionError("PDF dạng ảnh scan")

    _run(knowledge.process_document(lambda: session, _FakeEmbedder(), 6, extractor=_boom))
    assert doc.status == "processing"  # trạng thái 'failed' được ghi bằng câu UPDATE (session giả không áp)
    assert [type(s).__name__ for s in session.executed] == ["Update"]
    assert session.executed[0].compile().params["error_message"] == "PDF dạng ảnh scan"
