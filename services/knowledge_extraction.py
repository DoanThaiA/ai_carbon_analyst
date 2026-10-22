"""Trích chữ từ tài liệu tri thức admin tải lên (.pdf có lớp chữ / .docx) và chia đoạn để nhúng vector —
services/knowledge.py gọi. Toàn bộ là hàm thuần trên bytes/chuỗi (không MinIO/DB/LLM) để test được
(tests/test_knowledge_extraction.py).

- Định dạng xác định bằng magic bytes, KHÔNG tin MIME/đuôi file client khai báo.
- DOCX (python-docx): duyệt thân văn bản đúng thứ tự đoạn + bảng; style Heading/"Tiêu đề" mở đề mục mới; bảng →
  dòng "ô | ô | ô" (trafilatura/python-docx mặc định bỏ bảng, mà số liệu tham chiếu hay nằm trong bảng). Chặn zip
  bomb và file có macro trước khi parse.
- PDF (pdfplumber): theo từng trang — chữ ngoài vùng bảng + bảng dạng dòng; ghép lại đoạn văn từ các dòng; bỏ
  header/footer lặp và số trang. PDF dạng ảnh scan CHƯA hỗ trợ (chưa có OCR): trang gần như không có chữ mà có ảnh
  bị bỏ + cảnh báo; từ SCAN_FAIL_RATIO số trang trở lên như vậy → KnowledgeExtractionError.
- Mọi chữ chuẩn hoá Unicode NFC — PDF tiếng Việt hay ra dạng tổ hợp (NFD), làm hỏng tìm full-text và việc Jack
  đối chiếu số/chữ với nguồn.
"""
from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from typing import List, Optional

from services.chunking import CHUNK_MAX_CHARS, CHUNK_MIN_CHARS, chunk_text

KNOWLEDGE_MAX_BYTES = 20 * 1024 * 1024  # 20 MB — khớp frontend/src/lib/knowledgeUpload.ts
MAX_PAGES = 300
MAX_CHARS = 1_500_000
# Chống zip bomb (.docx là file zip): tổng dung lượng giải nén + số entry tối đa.
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCX_ENTRIES = 3000
# Trang PDF có ít hơn chừng này ký tự (không tính khoảng trắng) mà có ảnh = trang scan.
SCAN_PAGE_MIN_CHARS = 25
SCAN_FAIL_RATIO = 0.5
# Header/footer: dòng đầu/cuối trang lặp ở >= tỉ lệ này số trang (chỉ xét khi đủ số trang) bị bỏ.
HEADER_FOOTER_MIN_PAGES = 4
HEADER_FOOTER_RATIO = 0.5
HEADER_FOOTER_LINES = 2

MEDIA_PDF = "pdf"
MEDIA_DOCX = "docx"


class KnowledgeExtractionError(ValueError):
    """Tài liệu không trích được (sai định dạng, hỏng, có mật khẩu, scan, vượt giới hạn) — thông báo hiển thị cho admin."""


@dataclass
class ExtractedSection:
    heading: Optional[str]
    page_from: Optional[int]
    page_to: Optional[int]
    text: str

    def to_dict(self) -> dict:
        return {"heading": self.heading, "page_from": self.page_from, "page_to": self.page_to, "text": self.text}


@dataclass
class ExtractedDocument:
    media_type: str
    sections: List[ExtractedSection]
    page_count: Optional[int] = None
    warnings: List[str] = field(default_factory=list)
    title_hint: Optional[str] = None

    @property
    def content(self) -> str:
        parts = []
        for s in self.sections:
            parts.append(f"## {s.heading}\n\n{s.text}" if s.heading else s.text)
        return "\n\n".join(p for p in parts if p.strip())

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.sections)


@dataclass
class KnowledgeChunk:
    text: str
    heading: Optional[str]
    page_from: Optional[int]
    page_to: Optional[int]

    @property
    def meta(self) -> dict:
        return {"heading": self.heading, "page_from": self.page_from, "page_to": self.page_to}


# ── Chuẩn hoá chữ ────────────────────────────────────────────────────

_LIGATURES = str.maketrans({"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"})
_INVISIBLE_RE = re.compile(r"[­​‌‍⁠﻿]")
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b-\x1f\x7f]")
_SPACES_RE = re.compile(r"[ \t  -  　]+")


def normalize_text(text: str) -> str:
    """NFC + bỏ ký tự vô hình/điều khiển + gộp khoảng trắng; giữ xuống dòng (tối đa 1 dòng trống liên tiếp)."""
    text = unicodedata.normalize("NFC", text or "").translate(_LIGATURES)
    text = _INVISIBLE_RE.sub("", text).replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_RE.sub(" ", text)
    lines = [_SPACES_RE.sub(" ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


# ── Nhận diện định dạng ──────────────────────────────────────────────

def detect_media_type(data: bytes) -> str:
    """'pdf' | 'docx' theo magic bytes. .doc (OLE cũ), .docm (có macro) và mọi định dạng khác bị từ chối."""
    if b"%PDF-" in data[:1024]:
        return MEDIA_PDF
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise KnowledgeExtractionError(
            "File .doc (Word định dạng cũ) chưa được hỗ trợ — lưu lại dưới dạng .docx rồi tải lên lại."
        )
    if data[:4] == b"PK\x03\x04":
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            raise KnowledgeExtractionError("File Word bị hỏng (không đọc được cấu trúc .docx).")
        if "word/document.xml" in names:
            return MEDIA_DOCX
    raise KnowledgeExtractionError("Chỉ nhận PDF hoặc Word (.docx).")


def _check_docx_archive(data: bytes) -> None:
    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except zipfile.BadZipFile:
        raise KnowledgeExtractionError("File Word bị hỏng (không đọc được cấu trúc .docx).")
    if len(infos) > MAX_DOCX_ENTRIES or sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise KnowledgeExtractionError("File Word có cấu trúc bất thường (dung lượng giải nén quá lớn) — từ chối xử lý.")
    if any(i.filename.lower().endswith("vbaproject.bin") for i in infos):
        raise KnowledgeExtractionError("File Word có macro (.docm) — không nhận, lưu lại dưới dạng .docx không macro.")


# ── DOCX ─────────────────────────────────────────────────────────────

_HEADING_STYLE_RE = re.compile(r"^(heading|tiêu đề|title)\b", re.IGNORECASE)


def _table_lines(rows: List[List[str]]) -> List[str]:
    """Bảng → dòng "ô | ô | ô" — bỏ ô gộp lặp liền nhau (python-docx/pdfplumber trả lại cùng chữ cho ô gộp) và dòng rỗng."""
    lines = []
    for row in rows:
        cells: List[str] = []
        for cell in row:
            value = normalize_text(cell or "").replace("\n", " ")
            if not cells or value != cells[-1]:
                cells.append(value)
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_docx_archive(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise KnowledgeExtractionError("File Word bị hỏng hoặc không phải .docx hợp lệ.")

    sections: List[ExtractedSection] = []
    heading: Optional[str] = None
    blocks: List[str] = []

    def _flush() -> None:
        text = "\n\n".join(b for b in blocks if b.strip())
        if text.strip() or heading:
            sections.append(ExtractedSection(heading=heading, page_from=None, page_to=None, text=text))

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, document)
            text = normalize_text(para.text)
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if _HEADING_STYLE_RE.match(style) and len(text) <= 200:
                _flush()
                heading, blocks = text, []
            else:
                blocks.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            lines = _table_lines([[cell.text for cell in row.cells] for row in table.rows])
            if lines:
                blocks.append("\n".join(lines))
    _flush()
    sections = [s for s in sections if s.text.strip()]

    warnings: List[str] = []
    if body.findall(".//" + qn("w:txbxContent")):
        warnings.append("Tài liệu có hộp văn bản (text box) — chữ trong hộp không được đọc.")
    if len(document.inline_shapes):
        warnings.append(f"Tài liệu có {len(document.inline_shapes)} ảnh/biểu đồ — nội dung trong ảnh không được đọc.")
    if not sections:
        raise KnowledgeExtractionError("Không trích được chữ nào từ file Word (file rỗng hoặc chỉ có ảnh).")
    title = normalize_text(document.core_properties.title or "") or None
    return ExtractedDocument(media_type=MEDIA_DOCX, sections=sections, warnings=warnings, title_hint=title)


# ── PDF ──────────────────────────────────────────────────────────────

_PAGE_NUMBER_RE = re.compile(r"^(?:trang|page|tr\.)?\s*[-–—]?\s*\d{1,4}\s*(?:/\s*\d{1,4}|of\s+\d{1,4})?\s*[-–—]?$", re.IGNORECASE)
_SENTENCE_END_RE = re.compile(r"[.!?:;…\"”»)]$")
_LIST_START_RE = re.compile(r"^(?:[-–—•▪●◦*]\s|\(?\d{1,3}[.)]\s|\(?[a-zđ][.)]\s|(?:Điều|Article|Chương|Chapter|Mục|Section)\s+\d)", re.IGNORECASE)


def _norm_repeat_key(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip().lower())


def _repeated_margin_lines(pages_lines: List[List[str]]) -> set:
    """Dòng đầu/cuối trang (đã thay số bằng #) lặp ở >= HEADER_FOOTER_RATIO số trang → header/footer."""
    if len(pages_lines) < HEADER_FOOTER_MIN_PAGES:
        return set()
    counts: Counter = Counter()
    for lines in pages_lines:
        # Trang ít dòng: chỉ xét 1 dòng đầu/cuối, tránh coi thân bài là header/footer.
        k = max(1, min(HEADER_FOOTER_LINES, len(lines) // 3))
        margin = lines[:k] + lines[-k:]
        counts.update({_norm_repeat_key(line) for line in margin if line.strip()})
    threshold = HEADER_FOOTER_RATIO * len(pages_lines)
    return {key for key, n in counts.items() if n >= threshold}


def join_pdf_lines(lines: List[str]) -> str:
    """Dòng PDF (pdfplumber trả từng dòng hiển thị, không có dòng trống giữa đoạn) → đoạn văn cách nhau 1 dòng trống.
    Ngắt đoạn khi dòng trước kết thúc bằng dấu câu và ngắn hơn dòng điển hình, hoặc dòng sau là gạch đầu dòng/điều
    khoản; còn lại nối bằng khoảng trắng (từ tiếng Anh bị gạch nối cuối dòng thì nối liền)."""
    lines = [line for line in (line.strip() for line in lines) if line]
    if not lines:
        return ""
    typical = sorted(len(line) for line in lines)[int(len(lines) * 0.8)] if len(lines) >= 3 else max(len(line) for line in lines)
    paragraphs: List[str] = [lines[0]]
    for prev, line in zip(lines, lines[1:]):
        short_prev = len(prev) < 0.7 * typical
        if _LIST_START_RE.match(line) or (_SENTENCE_END_RE.search(prev) and short_prev) or (short_prev and line[:1].isupper()):
            paragraphs.append(line)
        elif re.search(r"[A-Za-z]-$", paragraphs[-1]) and line[:1].islower():
            paragraphs[-1] = paragraphs[-1][:-1] + line
        else:
            paragraphs[-1] += " " + line
    return "\n\n".join(paragraphs)


def _inside(obj: dict, bbox: tuple) -> bool:
    x0, top, x1, bottom = bbox
    cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
    return x0 <= cx <= x1 and top <= cy <= bottom


def extract_pdf(data: bytes) -> ExtractedDocument:
    import pdfplumber

    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception:
        raise KnowledgeExtractionError("Không mở được PDF (file hỏng hoặc có mật khẩu).")
    with pdf:
        page_count = len(pdf.pages)
        if page_count == 0:
            raise KnowledgeExtractionError("PDF không có trang nào.")
        if page_count > MAX_PAGES:
            raise KnowledgeExtractionError(f"PDF có {page_count} trang — vượt giới hạn {MAX_PAGES} trang, hãy tách nhỏ.")
        raw_pages: List[tuple] = []  # (số trang, dòng chữ ngoài bảng, dòng bảng)
        scan_pages: List[int] = []
        try:
            for no, page in enumerate(pdf.pages, start=1):
                tables = page.find_tables()
                bboxes = [t.bbox for t in tables]
                text_page = page.filter(lambda obj: not any(_inside(obj, b) for b in bboxes)) if bboxes else page
                text = normalize_text(text_page.extract_text() or "")
                table_blocks = [lines for lines in (_table_lines(t.extract()) for t in tables) if lines]
                if len(re.sub(r"\s", "", text)) < SCAN_PAGE_MIN_CHARS and not table_blocks:
                    if page.images:
                        scan_pages.append(no)
                    continue
                raw_pages.append((no, text.split("\n") if text else [], table_blocks))
        except KnowledgeExtractionError:
            raise
        except Exception:
            raise KnowledgeExtractionError("Lỗi khi đọc nội dung PDF (file có thể bị hỏng).")
        title = normalize_text(str((pdf.metadata or {}).get("Title") or "")) or None

    if scan_pages and len(scan_pages) >= SCAN_FAIL_RATIO * page_count:
        raise KnowledgeExtractionError(
            f"PDF dạng ảnh scan ({len(scan_pages)}/{page_count} trang không có lớp chữ) — hệ thống chưa hỗ trợ OCR; "
            "hãy dùng bản PDF có chữ hoặc .docx."
        )

    repeated = _repeated_margin_lines([lines for _, lines, _ in raw_pages])
    sections: List[ExtractedSection] = []
    for no, lines, table_blocks in raw_pages:
        kept = [
            line for line in lines
            if _norm_repeat_key(line) not in repeated and not _PAGE_NUMBER_RE.match(line.strip())
        ]
        parts = [join_pdf_lines(kept)] + ["\n".join(block) for block in table_blocks]
        text = "\n\n".join(p for p in parts if p.strip())
        if text:
            sections.append(ExtractedSection(heading=None, page_from=no, page_to=no, text=text))

    warnings: List[str] = []
    if scan_pages:
        warnings.append(f"Bỏ qua {len(scan_pages)} trang dạng ảnh scan (không có lớp chữ): " + _page_list(scan_pages))
    if not sections:
        raise KnowledgeExtractionError("Không trích được chữ nào từ PDF.")
    return ExtractedDocument(
        media_type=MEDIA_PDF, sections=sections, page_count=page_count, warnings=warnings, title_hint=title,
    )


def _page_list(pages: List[int], limit: int = 20) -> str:
    shown = ", ".join(str(p) for p in pages[:limit])
    return shown + (f"… (+{len(pages) - limit})" if len(pages) > limit else "")


def extract_document(data: bytes) -> ExtractedDocument:
    """Nhận diện định dạng + trích chữ. Raise KnowledgeExtractionError với thông báo cho admin."""
    if len(data) > KNOWLEDGE_MAX_BYTES:
        raise KnowledgeExtractionError(f"File vượt quá {KNOWLEDGE_MAX_BYTES // (1024 * 1024)}MB.")
    media_type = detect_media_type(data)
    doc = extract_pdf(data) if media_type == MEDIA_PDF else extract_docx(data)
    if doc.char_count > MAX_CHARS:
        raise KnowledgeExtractionError(
            f"Tài liệu quá dài ({doc.char_count:,} ký tự, tối đa {MAX_CHARS:,}) — hãy tách thành nhiều tài liệu."
        )
    return doc


# ── Chia đoạn ────────────────────────────────────────────────────────

def _is_table_block(block: str) -> bool:
    lines = block.split("\n")
    return len(lines) >= 2 and sum(" | " in line for line in lines) >= 0.6 * len(lines)


def _split_block(block: str, max_chars: int) -> List[str]:
    """Khối dài hơn max_chars: bảng → nhóm dòng, mỗi nhóm lặp lại dòng tiêu đề bảng; văn xuôi → chunk_text (theo câu)."""
    if len(block) <= max_chars:
        return [block]
    if not _is_table_block(block):
        return chunk_text(block.replace("\n", " "), max_chars=max_chars, min_chars=0)
    header, *rows = block.split("\n")
    groups: List[str] = []
    current: List[str] = []
    for row in rows:
        candidate = "\n".join([header, *current, row])
        if current and len(candidate) > max_chars:
            groups.append("\n".join([header, *current]))
            current = [row]
        else:
            current.append(row)
    if current:
        groups.append("\n".join([header, *current]))
    return groups


def chunk_document(
    sections: List[ExtractedSection], max_chars: int = CHUNK_MAX_CHARS, min_chars: int = CHUNK_MIN_CHARS,
) -> List[KnowledgeChunk]:
    """Gộp tham lam các khối (đoạn văn / bảng) thành chunk ~max_chars — KHÔNG vượt ranh giới đề mục (DOCX); PDF không
    có đề mục nên chunk có thể nối sang trang sau (page_from/page_to ghi đủ khoảng trang). Chunk cuối của 1 đề mục
    ngắn hơn min_chars được nối vào chunk trước nếu không vượt quá 1,2 × max_chars."""
    chunks: List[KnowledgeChunk] = []
    for sec in sections:
        for raw in re.split(r"\n\s*\n", sec.text):
            raw = raw.strip()
            if not raw:
                continue
            for block in _split_block(raw, max_chars):
                last = chunks[-1] if chunks else None
                if last and last.heading == sec.heading and len(last.text) + 2 + len(block) <= max_chars:
                    last.text = f"{last.text}\n\n{block}"
                    last.page_to = sec.page_to if sec.page_to is not None else last.page_to
                elif (
                    last and last.heading == sec.heading and len(block) < min_chars
                    and len(last.text) + 2 + len(block) <= int(max_chars * 1.2)
                ):
                    last.text = f"{last.text}\n\n{block}"
                    last.page_to = sec.page_to if sec.page_to is not None else last.page_to
                else:
                    chunks.append(KnowledgeChunk(text=block, heading=sec.heading, page_from=sec.page_from, page_to=sec.page_to))
    return chunks


def embedding_text(title: str, chunk: KnowledgeChunk) -> str:
    """Text đưa đi nhúng vector: thêm tiền tố tên tài liệu + đề mục (đoạn tách rời thường thiếu ngữ cảnh "văn bản nào,
    phần nào") — `chunks.content` vẫn lưu đoạn gốc không tiền tố."""
    prefix = f"Tài liệu: {title}" + (f" — {chunk.heading}" if chunk.heading else "")
    return f"{prefix}\n{chunk.text}"
